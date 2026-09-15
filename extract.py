# -*- coding: utf-8 -*-
"""Dosyadan metin ve 'belge adli sinyalleri' cikarma."""
import io
import re
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime


def _parse_dt(s):
    if not s:
        return None
    s = s.strip().replace("Z", "")
    for fmt in ("%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M"):
        try:
            return datetime.strptime(s[:19], fmt)
        except ValueError:
            continue
    return None


# ---------------------------------------------------------------- DOCX ------

def from_docx(data):
    import docx

    doc = docx.Document(io.BytesIO(data))
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.append(cell.text)
    text = "\n\n".join(c.strip() for c in chunks if c and c.strip())

    notes, delta = [], 0
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            names = set(z.namelist())
            app = core = None
            if "docProps/app.xml" in names:
                app = ET.fromstring(z.read("docProps/app.xml"))
            if "docProps/core.xml" in names:
                core = ET.fromstring(z.read("docProps/core.xml"))

            def tag(root, local):
                if root is None:
                    return None
                for el in root.iter():
                    if el.tag.split("}")[-1] == local:
                        return (el.text or "").strip()
                return None

            words = len(text.split())
            total_time = tag(app, "TotalTime")
            application = tag(app, "Application")
            revision = tag(core, "revision")
            created = _parse_dt(tag(core, "created"))
            modified = _parse_dt(tag(core, "modified"))

            if application:
                notes.append({"label": "Üreten uygulama", "value": application, "flag": "info"})

            if total_time is not None and total_time.isdigit():
                tt = int(total_time)
                notes.append({"label": "Toplam düzenleme süresi",
                              "value": "%d dakika" % tt,
                              "flag": "warn" if (tt <= 3 and words > 400) else "info"})
                if tt <= 1 and words > 400:
                    delta += 7
                    notes.append({"label": "Şüphe", "flag": "warn",
                                  "value": "%d kelimelik belge ~%d dk'da oluşmuş — "
                                           "metnin dışarıdan yapıştırıldığına işaret eder." % (words, tt)})
                elif tt <= 3 and words > 800:
                    delta += 4

            if revision and revision.isdigit():
                rev = int(revision)
                notes.append({"label": "Kayıt (revizyon) sayısı", "value": str(rev),
                              "flag": "warn" if rev <= 2 and words > 400 else "info"})
                if rev <= 1 and words > 400:
                    delta += 3

            if created and modified:
                gap = (modified - created).total_seconds()
                notes.append({"label": "Oluşturma → son kayıt",
                              "value": ("%d dk" % round(gap / 60)) if gap >= 60 else "%d sn" % round(gap),
                              "flag": "warn" if gap < 120 and words > 400 else "info"})
                if gap < 60 and words > 400:
                    delta += 3
    except Exception:
        pass

    return text, {"notes": notes, "score_delta": min(delta, 10)}


# ----------------------------------------------------------------- PDF ------

SUSPECT_PRODUCERS = {
    "skia/pdf": "Tarayıcıdan (Chrome) yazdırılmış — web sayfası veya sohbet ekranı kaynaklı olabilir.",
    "wkhtmltopdf": "HTML'den otomatik üretilmiş.",
    "weasyprint": "HTML'den otomatik üretilmiş.",
    "reportlab": "Bir betik/program tarafından üretilmiş.",
    "puppeteer": "Otomatik tarayıcı ile üretilmiş.",
}


def from_pdf(data):
    from pdfminer.high_level import extract_text
    from pdfminer.pdfparser import PDFParser
    from pdfminer.pdfdocument import PDFDocument

    text = extract_text(io.BytesIO(data)) or ""
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"(?<=[a-zçğıöşü])-\n(?=[a-zçğıöşü])", "", text)  # satir sonu tireleme

    notes, delta = [], 0
    try:
        parser = PDFParser(io.BytesIO(data))
        doc = PDFDocument(parser)
        info = doc.info[0] if doc.info else {}

        def val(k):
            v = info.get(k)
            if isinstance(v, bytes):
                try:
                    return v.decode("utf-16" if v[:2] in (b"\xfe\xff", b"\xff\xfe") else "utf-8",
                                    errors="replace").strip()
                except Exception:
                    return v.decode("latin-1", errors="replace").strip()
            return str(v).strip() if v else None

        producer, creator = val("Producer"), val("Creator")
        if producer:
            notes.append({"label": "PDF üreticisi", "value": producer, "flag": "info"})
        if creator and creator != producer:
            notes.append({"label": "Oluşturan program", "value": creator, "flag": "info"})

        blob = " ".join(filter(None, [producer, creator])).lower()
        for key, msg in SUSPECT_PRODUCERS.items():
            if key in blob:
                notes.append({"label": "Dikkat", "value": msg, "flag": "warn"})
                delta += 3
                break

        raw_c, raw_m = val("CreationDate"), val("ModDate")
        if raw_c and raw_m and raw_c[:16] == raw_m[:16]:
            notes.append({"label": "Oluşturma = değiştirme zamanı",
                          "value": "Tek seferde üretilmiş (düzenleme geçmişi yok)", "flag": "info"})
    except Exception:
        pass

    if len(text.split()) < 25:
        raise ValueError("PDF'den metin çıkarılamadı. Belge taranmış görüntü olabilir; "
                         "bu durumda önce OCR gerekir.")

    return text, {"notes": notes, "score_delta": min(delta, 10)}


# ----------------------------------------------------------------- Giris ----

def extract(filename, data):
    name = (filename or "").lower()
    if name.endswith(".pdf") or data[:5] == b"%PDF-":
        return from_pdf(data)
    if name.endswith(".docx") or data[:2] == b"PK":
        return from_docx(data)
    if name.endswith((".txt", ".md", ".rtf", ".csv")):
        for enc in ("utf-8", "utf-16", "cp1254", "latin-1"):
            try:
                return data.decode(enc), {"notes": [], "score_delta": 0}
            except UnicodeDecodeError:
                continue
    if name.endswith(".doc"):
        raise ValueError("Eski .doc biçimi desteklenmiyor. Word'de “Farklı Kaydet → .docx” "
                         "yapıp tekrar deneyin.")
    raise ValueError("Desteklenmeyen dosya türü. PDF, DOCX, TXT veya MD yükleyin.")
